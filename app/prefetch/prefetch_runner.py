"""
사전 쿼리 실행기 — Qings 엑셀에서 SN 추출 후 Superset 쿼리를 사전 실행합니다.
"""
import asyncio
import logging
import os
from datetime import datetime
from typing import Optional

from app.config import settings
from app.prefetch.cache_manager import is_cached

logger = logging.getLogger(__name__)

# ── 실행 상태 ─────────────────────────────────────────────────────────────────
_status: dict = {
    "running": False,
    "last_run": None,
    "last_result": None,
    "error": None,
}


def get_status() -> dict:
    return dict(_status)


# ── SN 추출 ───────────────────────────────────────────────────────────────────
def _iter_excel_rows(excel_path: str):
    """xls/xlsx 모두 지원하는 행 이터레이터. (headers, row_iter) 반환."""
    ext = str(excel_path).lower()

    if ext.endswith(".xls"):
        # 1) xlrd (진짜 BIFF 포맷)
        try:
            import xlrd
            wb = xlrd.open_workbook(excel_path)
            ws = wb.sheet_by_index(0)
            headers = [str(ws.cell_value(0, c)).strip() for c in range(ws.ncols)]
            def _xlrd_rows():
                for r in range(1, ws.nrows):
                    yield tuple(ws.cell_value(r, c) for c in range(ws.ncols))
            return headers, _xlrd_rows()
        except Exception as e:
            logger.warning(f"[Prefetch] xlrd 실패({e}), openpyxl 시도...")

        # 2) openpyxl (xlsx 확장자를 .xls로 저장한 경우)
        try:
            import openpyxl
            wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
            ws = wb.active
            rows = ws.iter_rows(values_only=True)
            header_row = next(rows)
            headers = [str(v).strip() if v is not None else "" for v in header_row]
            return headers, rows
        except Exception as e:
            logger.warning(f"[Prefetch] openpyxl 실패({e}), HTML 파싱 시도...")

        # 3) HTML 테이블 (한국 기업 시스템에서 흔한 HTML-as-XLS)
        import html.parser, pathlib

        class _TableParser(html.parser.HTMLParser):
            def __init__(self):
                super().__init__()
                self.rows: list[list[str]] = []
                self._row: list[str] = []
                self._cell = False
                self._data: list[str] = []
            def handle_starttag(self, tag, attrs):
                if tag in ("tr",):
                    self._row = []
                elif tag in ("td", "th"):
                    self._cell = True
                    self._data = []
            def handle_endtag(self, tag):
                if tag in ("td", "th"):
                    self._row.append("".join(self._data).strip())
                    self._cell = False
                elif tag == "tr" and self._row:
                    self.rows.append(self._row)
            def handle_data(self, data):
                if self._cell:
                    self._data.append(data)

        raw = pathlib.Path(excel_path).read_bytes()
        for enc in ("utf-8", "euc-kr", "cp949"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            text = raw.decode("utf-8", errors="replace")

        parser = _TableParser()
        parser.feed(text)
        if not parser.rows:
            raise ValueError("HTML 파싱 결과 없음")

        # 셀 수가 가장 많은 행을 헤더로 사용 (제목행은 보통 셀 1개)
        header_idx = max(range(len(parser.rows)), key=lambda i: len(parser.rows[i]))
        logger.info(f"[Prefetch] HTML 헤더 행: {header_idx}번째 ({len(parser.rows[header_idx])}열)")

        headers = [c.strip() for c in parser.rows[header_idx]]
        data_rows = parser.rows[header_idx + 1:]
        def _html_rows():
            for r in data_rows:
                yield tuple(r)
        return headers, _html_rows()

    else:
        import openpyxl
        wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        header_row = next(rows)
        headers = [str(v).strip() if v is not None else "" for v in header_row]
        return headers, rows


def _load_filters() -> list[tuple[str, set[str]]]:
    """filters.json에서 (column_name, allowed_values) 목록을 반환합니다."""
    import json
    from pathlib import Path
    filters_path = Path(__file__).parent.parent.parent / "filters.json"
    if not filters_path.exists():
        return []
    try:
        data = json.loads(filters_path.read_text(encoding="utf-8"))
        result = []
        for item in data.get("filters", {}).values():
            col = item.get("column", "").strip()
            vals = {str(v).strip() for v in item.get("values", []) if str(v).strip()}
            if col and vals:
                result.append((col, vals))
        return result
    except Exception as e:
        logger.warning(f"[Prefetch] filters.json 로드 실패: {e}")
        return []


def extract_sns_from_excel(excel_path: str, sn_column: str | None = None) -> list[str]:
    """xlsx/xls에서 SN 목록을 추출합니다.

    filters.json 조건(AND)으로 행 필터링 후 SN 추출, 중복 제거.
    """
    if sn_column is None:
        sn_column = settings.QINGS_SN_COLUMN

    headers, row_iter = _iter_excel_rows(excel_path)

    if sn_column in headers:
        sn_idx = headers.index(sn_column)
        logger.info(f"[Prefetch] SN 열 '{sn_column}' → index {sn_idx}")
    elif settings.QINGS_SN_COL_IDX >= 0:
        sn_idx = settings.QINGS_SN_COL_IDX
        logger.warning(
            f"[Prefetch] SN 열 '{sn_column}' 없음 → 설정된 열 인덱스 {sn_idx}({chr(65+sn_idx)}열) 사용. "
            f"헤더: {headers[:10]}"
        )
    else:
        logger.warning(f"[Prefetch] SN 열 '{sn_column}' 없음. 헤더: {headers[:10]}")
        return []

    # filters.json 로드 및 열 인덱스 매핑
    filter_specs = _load_filters()
    filter_idxs: list[tuple[int, set[str]]] = []
    for col_name, allowed in filter_specs:
        if col_name in headers:
            filter_idxs.append((headers.index(col_name), allowed))
        else:
            logger.warning(f"[Prefetch] 필터 열 '{col_name}' 없음 — 해당 조건 무시")

    if filter_idxs:
        logger.info(f"[Prefetch] 필터 {len(filter_idxs)}개 적용 (AND)")
    else:
        logger.warning("[Prefetch] 적용할 필터 없음 — 전체 SN 추출")

    seen: set[str] = set()
    sns:  list[str] = []
    skipped = 0

    for row in row_iter:
        # AND 필터: 모든 조건 만족해야 통과
        passed = True
        for idx, allowed in filter_idxs:
            cell_val = str(row[idx]).strip() if idx < len(row) and row[idx] is not None else ""
            if cell_val not in allowed:
                passed = False
                break
        if not passed:
            skipped += 1
            continue

        val = row[sn_idx] if sn_idx < len(row) else None
        if val:
            sn = str(val).strip().upper()
            if sn and sn not in seen:
                seen.add(sn)
                sns.append(sn)

    logger.info(f"[Prefetch] SN {len(sns)}개 추출 (필터 통과 / {skipped}개 제외)")
    return sns


# ── SN 목록으로 사전 쿼리 ────────────────────────────────────────────────────
async def run_prefetch_for_sns(sns: list[str]) -> dict:
    """주어진 SN 목록에 대해 캐시 없는 항목만 Superset 쿼리를 실행합니다."""
    from app.scraper.query_runner import query_runner

    if _status["running"]:
        logger.warning("[Prefetch] 이미 실행 중입니다.")
        return {"status": "already_running"}

    _status["running"] = True
    _status["error"] = None

    uncached = [sn for sn in sns if not is_cached(sn)]
    cached_count = len(sns) - len(uncached)
    result = {
        "total": len(sns),
        "cached": cached_count,
        "queried": 0,
        "failed": 0,
        "failed_sns": [],
        "sn_results": [],   # [{"sn": ..., "status": ..., "detail": ...}]
    }

    # 캐시된 SN은 미리 결과에 추가
    for sn in sns:
        if is_cached(sn):
            result["sn_results"].append({"sn": sn, "status": "캐시", "detail": ""})

    logger.info(
        f"[Prefetch] 시작 — 전체 {len(sns)}개 | 캐시 {cached_count}개 | "
        f"신규 쿼리 {len(uncached)}개"
    )

    try:
        for i, sn in enumerate(uncached, 1):
            logger.info(f"[Prefetch] [{i}/{len(uncached)}] {sn} 쿼리 중...")
            try:
                qr = await query_runner.run(sn)
                if qr.success:
                    if qr.csv_path:
                        status, detail = "성공", ""
                    else:
                        status, detail = "데이터없음", qr.error or ""
                    result["queried"] += 1
                    logger.info(f"[Prefetch] [{sn}] 완료 → {qr.csv_path}")
                else:
                    status, detail = "실패", qr.error or ""
                    result["failed"] += 1
                    result["failed_sns"].append(sn)
                    logger.warning(f"[Prefetch] [{sn}] 실패: {qr.error}")
            except Exception as e:
                status, detail = "오류", str(e)
                result["failed"] += 1
                result["failed_sns"].append(sn)
                logger.error(f"[Prefetch] [{sn}] 오류: {e}")

            result["sn_results"].append({"sn": sn, "status": status, "detail": detail})
            await asyncio.sleep(3)  # 서버 부하 분산

        _status["last_result"] = result
        logger.info(
            f"[Prefetch] 완료 — 쿼리 성공 {result['queried']}개 | "
            f"실패 {result['failed']}개"
        )
        return result

    except Exception as e:
        _status["error"] = str(e)
        logger.error(f"[Prefetch] 실행 오류: {e}", exc_info=True)
        return {"status": "error", "error": str(e)}
    finally:
        _status["running"] = False
        _status["last_run"] = datetime.now().isoformat()


# ── 전체 파이프라인 (Qings → 엑셀 → SN 추출 → 쿼리) ─────────────────────────
async def run_daily_prefetch() -> dict:
    """Qings에서 SN 수집 후 사전 쿼리 전체 파이프라인을 실행합니다."""
    from app.prefetch.qings_scraper import scrape_qings_excel

    if _status["running"]:
        return {"status": "already_running"}

    logger.info("[Prefetch] 일일 사전 쿼리 시작")

    # 1. Qings 엑셀 다운로드
    excel_path = await scrape_qings_excel(settings.CSV_DOWNLOAD_PATH)
    if not excel_path:
        _status["error"] = "Qings 엑셀 다운로드 실패"
        return {"status": "error", "error": _status["error"]}

    # 2. SN 추출
    sns = extract_sns_from_excel(excel_path)
    if not sns:
        return {"status": "error", "error": f"SN 추출 실패 — 열 이름 확인 필요: {settings.QINGS_SN_COLUMN}"}

    # 3. 사전 쿼리 실행
    return await run_prefetch_for_sns(sns)
