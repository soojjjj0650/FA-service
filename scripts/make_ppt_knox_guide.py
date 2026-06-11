"""
Knox Teams 메신저 만들기 - PPT 가이드 생성
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

# ── 색상 ──────────────────────────────────────────────────────────────────────
NAVY    = RGBColor(0x0f, 0x17, 0x2a)
BLUE    = RGBColor(0x1e, 0x3a, 0x6e)
MBLUE   = RGBColor(0x2d, 0x5a, 0x9e)
LBLUE   = RGBColor(0x60, 0xa5, 0xfa)
ACCENT  = RGBColor(0x38, 0xbd, 0xf8)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
GRAY    = RGBColor(0xf1, 0xf5, 0xf9)
DGRAY   = RGBColor(0x64, 0x74, 0x8b)
GREEN   = RGBColor(0x10, 0xb9, 0x81)
ORANGE  = RGBColor(0xf5, 0x9e, 0x0b)
RED     = RGBColor(0xef, 0x44, 0x44)
CODE_BG = RGBColor(0x1e, 0x29, 0x3b)
CODE_FG = RGBColor(0x93, 0xc5, 0xfd)
STR_FG  = RGBColor(0x86, 0xef, 0xac)
KEY_FG  = RGBColor(0xfb, 0xcf, 0x89)

W = Inches(13.33)
H = Inches(7.5)

prs = Presentation()
prs.slide_width  = W
prs.slide_height = H
BLANK = prs.slide_layouts[6]


# ── 헬퍼 ──────────────────────────────────────────────────────────────────────
def rect(sl, x, y, w, h, fill=None, line_color=None, line_pt=1):
    s = sl.shapes.add_shape(1, x, y, w, h)
    s.line.fill.background()
    if fill:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    else:
        s.fill.background()
    if line_color:
        s.line.color.rgb = line_color; s.line.width = Pt(line_pt)
    else:
        s.line.fill.background()
    return s


def txt(sl, text, x, y, w, h, size=18, bold=False, color=WHITE,
        align=PP_ALIGN.LEFT, italic=False, wrap=True):
    t = sl.shapes.add_textbox(x, y, w, h)
    tf = t.text_frame; tf.word_wrap = wrap
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold
    r.font.color.rgb = color; r.font.italic = italic
    return t


def header(sl, step_label, title, subtitle=None):
    rect(sl, 0, 0, W, H, fill=NAVY)
    rect(sl, 0, 0, W, Inches(0.1), fill=ACCENT)
    if step_label:
        txt(sl, step_label, Inches(0.5), Inches(0.18), Inches(3), Inches(0.5),
            size=12, bold=True, color=ACCENT)
    txt(sl, title, Inches(0.5), Inches(0.55) if step_label else Inches(0.25),
        Inches(12), Inches(0.75), size=26, bold=True, color=WHITE)
    if subtitle:
        txt(sl, subtitle, Inches(0.5), Inches(1.18), Inches(12), Inches(0.45),
            size=13, color=LBLUE, italic=True)
    rect(sl, Inches(0.4), Inches(1.5), W - Inches(0.8), Pt(1.5), fill=ACCENT)


def photo_box(sl, x, y, w, h, label="캡쳐 필요"):
    rect(sl, x, y, w, h, fill=RGBColor(0x12, 0x22, 0x3a), line_color=LBLUE, line_pt=1.5)
    txt(sl, f"📷   {label}", x, y + h/2 - Inches(0.35), w, Inches(0.7),
        size=15, bold=True, color=LBLUE, align=PP_ALIGN.CENTER)


def path_box(sl, x, y, w, path_text):
    rect(sl, x, y, w, Inches(0.5), fill=BLUE)
    rect(sl, x, y, Inches(0.06), Inches(0.5), fill=ACCENT)
    txt(sl, path_text, x + Inches(0.15), y + Inches(0.06), w - Inches(0.2), Inches(0.4),
        size=12, bold=True, color=WHITE)


def info_row(sl, x, y, w, label, value, label_color=ACCENT):
    txt(sl, label, x, y, Inches(2.5), Inches(0.38),
        size=12, bold=True, color=label_color)
    txt(sl, value, x + Inches(2.6), y, w - Inches(2.6), Inches(0.38),
        size=12, color=WHITE)


def note_box(sl, x, y, w, text, color=ORANGE):
    rect(sl, x, y, w, Inches(0.5), fill=CODE_BG)
    rect(sl, x, y, Inches(0.06), Inches(0.5), fill=color)
    txt(sl, text, x + Inches(0.15), y + Inches(0.06),
        w - Inches(0.2), Inches(0.4), size=11, color=WHITE)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 1 — 표지
# ══════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
rect(sl, 0, 0, W, H, fill=NAVY)
rect(sl, 0, 0, Inches(0.3), H, fill=ACCENT)
rect(sl, Inches(9.2), Inches(1.0), Inches(4.8), Inches(6.0), fill=BLUE)
rect(sl, Inches(10.0), Inches(2.0), Inches(4.0), Inches(4.5), fill=RGBColor(0x1a, 0x32, 0x6b))

txt(sl, "💬", Inches(0.7), Inches(1.3), Inches(1.8), Inches(1.4), size=58, color=WHITE)
txt(sl, "Knox Teams\n메신저 만들기",
    Inches(0.7), Inches(2.5), Inches(8.2), Inches(2.3),
    size=42, bold=True, color=WHITE)
txt(sl, "봇 등록  ·  방화벽  ·  서버 등록  ·  API 연결",
    Inches(0.7), Inches(4.9), Inches(8.2), Inches(0.55),
    size=15, color=LBLUE, italic=True)
txt(sl, "Knox C&C Suite Developer Guide",
    Inches(0.7), Inches(6.1), Inches(6.5), Inches(0.5),
    size=13, color=DGRAY)
txt(sl, "http://developers.samsung.net/static/knoxcenter/main.html#",
    Inches(0.7), Inches(6.6), Inches(8.0), Inches(0.45),
    size=11, color=LBLUE, italic=True)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 2 — 전체 흐름 개요
# ══════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
header(sl, None, "전체 구성", "이 순서대로 진행합니다")

steps = [
    ("01", "챗봇 아이디\n생성",    "Knox Portal에서\n봇 계정 신청"),
    ("02", "방화벽\n연결",         "서버 포트 개방\n(IN/OUT)"),
    ("03", "Knox 서버\n등록",      "System ID &\nToken 발급"),
    ("04", "API 연결\n챗봇 방 생성", "Device 등록\n대화방 생성"),
    ("05", "FAQ",                  "문의처 &\nKnox Support"),
]

bw = Inches(2.3)
bh = Inches(3.8)
gap = Inches(0.17)
total = len(steps) * bw + (len(steps) - 1) * gap
sx = (W - total) / 2
by = Inches(1.7)

for i, (num, title, desc) in enumerate(steps):
    bx = sx + i * (bw + gap)
    c = NAVY if i % 2 == 0 else BLUE
    rect(sl, bx, by, bw, bh, fill=c)
    rect(sl, bx, by, bw, Inches(0.08), fill=ACCENT)
    txt(sl, num, bx, by + Inches(0.12), bw, Inches(0.5),
        size=22, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
    txt(sl, title, bx, by + Inches(0.65), bw, Inches(1.1),
        size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    txt(sl, desc, bx, by + Inches(1.8), bw, Inches(1.8),
        size=11, color=LBLUE, align=PP_ALIGN.CENTER)
    if i < len(steps) - 1:
        txt(sl, "›", bx + bw, by + Inches(1.5), gap + Inches(0.08), Inches(0.6),
            size=18, color=ACCENT, align=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 3 — STEP 01: 챗봇 아이디 생성하기
# ══════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
header(sl, "STEP 01", "챗봇 아이디 생성하기",
       "Knox Portal에서 봇 계정을 신청합니다. 결재 후 챗봇 아이디가 생성됩니다.")

# 경로 안내
path_box(sl, Inches(0.5), Inches(1.7),
         W - Inches(1.0),
         "Home  →  My Interface  →  Bot  →  봇관리  →  봇계정신청")

# 결재 정보
rect(sl, Inches(0.5), Inches(2.35), Inches(3.8), Inches(2.0), fill=BLUE)
rect(sl, Inches(0.5), Inches(2.35), Inches(3.8), Inches(0.08), fill=ACCENT)
txt(sl, "  결재 정보",
    Inches(0.5), Inches(2.38), Inches(3.8), Inches(0.42),
    size=13, bold=True, color=WHITE)
for i, item in enumerate(["파트장  결재", "경영지원  합의"]):
    txt(sl, f"•  {item}",
        Inches(0.7), Inches(2.9) + i * Inches(0.48), Inches(3.4), Inches(0.45),
        size=13, color=LBLUE)

# 캡쳐 박스
photo_box(sl, Inches(4.5), Inches(2.35), Inches(8.5), Inches(4.7),
          "봇계정신청 화면 캡쳐")

note_box(sl, Inches(0.5), Inches(4.5), Inches(3.8),
         "💡 결재 완료 후 아이디 생성 확인", color=GREEN)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 4 — STEP 02: 방화벽 개념도 (Knox 수신/발신 서버 구분, 밝은 배경)
# ══════════════════════════════════════════════════════════════════════════════
_LBGC  = RGBColor(0xf8, 0xf9, 0xfc)
_SLATE = RGBColor(0x47, 0x55, 0x69)

sl = prs.slides.add_slide(BLANK)
rect(sl, 0, 0, W, H, fill=_LBGC)
rect(sl, 0, 0, W, Inches(0.1), fill=NAVY)
txt(sl, "STEP 02", Inches(0.5), Inches(0.18), Inches(3), Inches(0.5),
    size=12, bold=True, color=MBLUE)
txt(sl, "방화벽 연결하기 — 개념도",
    Inches(0.5), Inches(0.55), Inches(12), Inches(0.7),
    size=26, bold=True, color=NAVY)
txt(sl, "Knox 서버는 수신용(API)과 발신용(Webhook) IP가 다릅니다 — 두 방향 모두 방화벽 정책 등록 필요",
    Inches(0.5), Inches(1.18), Inches(12), Inches(0.42),
    size=13, color=_SLATE, italic=True)
rect(sl, Inches(0.4), Inches(1.55), W - Inches(0.8), Pt(1.5), fill=NAVY)

_bw   = Inches(2.9)
_bh   = Inches(1.9)
_bgap = Inches(0.9)
_btot = 3 * _bw + 2 * _bgap
_bsx  = (W - _btot) / 2

# ── Row 1: 발신 (Outbound) — 내 서버 → Knox 수신 서버 ────────────────────────
_r1y = Inches(1.68)
rect(sl, _bsx, _r1y, _btot, Inches(0.3), fill=NAVY)
txt(sl, "  ▲  발신 (Outbound) — 내 서버  →  Knox 수신 서버 (API)  |  Port 80, 443",
    _bsx, _r1y, _btot, Inches(0.3), size=11, bold=True, color=WHITE)

_by1 = _r1y + Inches(0.34)
_fwx = _bsx + _bw + _bgap
_knx = _fwx + _bw + _bgap

# 내 서버
rect(sl, _bsx, _by1, _bw, _bh, fill=NAVY)
rect(sl, _bsx, _by1, _bw, Inches(0.08), fill=ACCENT)
txt(sl, "🖥️",  _bsx, _by1 + Inches(0.1),  _bw, Inches(0.65), size=26, color=WHITE, align=PP_ALIGN.CENTER)
txt(sl, "내 서버", _bsx, _by1 + Inches(0.73), _bw, Inches(0.4),  size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
txt(sl, "My Server  Port 80", _bsx, _by1 + Inches(1.18), _bw, Inches(0.32), size=10, color=LBLUE, align=PP_ALIGN.CENTER)

txt(sl, "Port\n80/443", _bsx + _bw, _by1 + Inches(0.12), _bgap, Inches(0.5), size=9, color=MBLUE, align=PP_ALIGN.CENTER)
txt(sl, "────▶",       _bsx + _bw, _by1 + Inches(0.6),  _bgap, Inches(0.5), size=20, bold=True, color=MBLUE, align=PP_ALIGN.CENTER)

# 방화벽
rect(sl, _fwx, _by1, _bw, _bh, fill=_SLATE)
rect(sl, _fwx, _by1, _bw, Inches(0.08), fill=DGRAY)
txt(sl, "🛡️",    _fwx, _by1 + Inches(0.1),  _bw, Inches(0.65), size=26, color=WHITE, align=PP_ALIGN.CENTER)
txt(sl, "방화벽",  _fwx, _by1 + Inches(0.73), _bw, Inches(0.4),  size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
txt(sl, "Firewall", _fwx, _by1 + Inches(1.18), _bw, Inches(0.32), size=10, color=RGBColor(0xcb, 0xd5, 0xe1), align=PP_ALIGN.CENTER)

txt(sl, "────▶", _fwx + _bw, _by1 + Inches(0.6), _bgap, Inches(0.5), size=20, bold=True, color=MBLUE, align=PP_ALIGN.CENTER)

# Knox 수신 서버 (API 요청을 받는 쪽)
rect(sl, _knx, _by1, _bw, _bh, fill=MBLUE)
rect(sl, _knx, _by1, _bw, Inches(0.08), fill=ACCENT)
txt(sl, "📡",           _knx, _by1 + Inches(0.1),  _bw, Inches(0.65), size=26, color=WHITE, align=PP_ALIGN.CENTER)
txt(sl, "Knox 수신 서버", _knx, _by1 + Inches(0.73), _bw, Inches(0.4),  size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
txt(sl, "openapi.samsung.net\nAPI 요청 수신",  _knx, _by1 + Inches(1.18), _bw, Inches(0.62), size=9.5, color=LBLUE, align=PP_ALIGN.CENTER)

# ── Row 2: 수신 (Inbound) — Knox 발신 서버 → 내 서버 ─────────────────────────
_r2y = _by1 + _bh + Inches(0.35)
rect(sl, _bsx, _r2y, _btot, Inches(0.3), fill=_SLATE)
txt(sl, "  ▼  수신 (Inbound) — Knox 발신 서버 (Webhook)  →  내 서버  |  Port 80  ⚠ 반드시 허용",
    _bsx, _r2y, _btot, Inches(0.3), size=11, bold=True, color=WHITE)

_by2 = _r2y + Inches(0.34)

# 내 서버 (수신 목적지)
rect(sl, _bsx, _by2, _bw, _bh, fill=NAVY)
rect(sl, _bsx, _by2, _bw, Inches(0.08), fill=ACCENT)
txt(sl, "🖥️",  _bsx, _by2 + Inches(0.1),  _bw, Inches(0.65), size=26, color=WHITE, align=PP_ALIGN.CENTER)
txt(sl, "내 서버", _bsx, _by2 + Inches(0.73), _bw, Inches(0.4),  size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
txt(sl, "My Server  Port 80", _bsx, _by2 + Inches(1.18), _bw, Inches(0.32), size=10, color=LBLUE, align=PP_ALIGN.CENTER)

txt(sl, "Port\n80",  _bsx + _bw, _by2 + Inches(0.12), _bgap, Inches(0.5), size=9, color=_SLATE, align=PP_ALIGN.CENTER)
txt(sl, "◀────", _bsx + _bw, _by2 + Inches(0.6),  _bgap, Inches(0.5), size=20, bold=True, color=_SLATE, align=PP_ALIGN.CENTER)

# 방화벽 (수신)
rect(sl, _fwx, _by2, _bw, _bh, fill=_SLATE)
rect(sl, _fwx, _by2, _bw, Inches(0.08), fill=DGRAY)
txt(sl, "🛡️",    _fwx, _by2 + Inches(0.1),  _bw, Inches(0.65), size=26, color=WHITE, align=PP_ALIGN.CENTER)
txt(sl, "방화벽",  _fwx, _by2 + Inches(0.73), _bw, Inches(0.4),  size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
txt(sl, "Firewall", _fwx, _by2 + Inches(1.18), _bw, Inches(0.32), size=10, color=RGBColor(0xcb, 0xd5, 0xe1), align=PP_ALIGN.CENTER)

txt(sl, "◀────", _fwx + _bw, _by2 + Inches(0.6), _bgap, Inches(0.5), size=20, bold=True, color=_SLATE, align=PP_ALIGN.CENTER)

# Knox 발신 서버 (Webhook을 보내는 쪽)
rect(sl, _knx, _by2, _bw, _bh, fill=BLUE)
rect(sl, _knx, _by2, _bw, Inches(0.08), fill=MBLUE)
txt(sl, "📡",           _knx, _by2 + Inches(0.1),  _bw, Inches(0.65), size=26, color=WHITE, align=PP_ALIGN.CENTER)
txt(sl, "Knox 발신 서버", _knx, _by2 + Inches(0.73), _bw, Inches(0.4),  size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
txt(sl, "Webhook POST 전송\n메시지를 내 서버로 발신", _knx, _by2 + Inches(1.18), _bw, Inches(0.62), size=9.5, color=LBLUE, align=PP_ALIGN.CENTER)

# 하단 안내
_noty = _by2 + _bh + Inches(0.22)
rect(sl, _bsx, _noty, _btot, Inches(0.52), fill=RGBColor(0xe4, 0xe9, 0xf4))
rect(sl, _bsx, _noty, Inches(0.06), Inches(0.52), fill=NAVY)
txt(sl, "  💡  Knox 서버 IP는 방향(수신/발신)과 환경(스테이지/운영)에 따라 다릅니다 — 다음 슬라이드에서 상세 확인",
    _bsx + Inches(0.1), _noty + Inches(0.06), _btot - Inches(0.2), Inches(0.42),
    size=11, color=NAVY)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 5 — STEP 02: 방화벽 정책 상세 (Stage + Production, 빨간색 제거)
# ══════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
rect(sl, 0, 0, W, H, fill=_LBGC)
rect(sl, 0, 0, W, Inches(0.1), fill=NAVY)
txt(sl, "STEP 02", Inches(0.5), Inches(0.18), Inches(3), Inches(0.5),
    size=12, bold=True, color=MBLUE)
txt(sl, "방화벽 정책 상세",
    Inches(0.5), Inches(0.55), Inches(12), Inches(0.7),
    size=26, bold=True, color=NAVY)
txt(sl, "스테이지(개발)와 운영 서버의 IP/포트 정보를 방화벽 담당자에게 전달합니다",
    Inches(0.5), Inches(1.18), Inches(12), Inches(0.42),
    size=13, color=_SLATE, italic=True)
rect(sl, Inches(0.4), Inches(1.55), W - Inches(0.8), Pt(1.5), fill=NAVY)

def fw_panel(sl, ox, title, hdr_color, out_ip, out_host, in_ips):
    pw = Inches(6.0)
    rect(sl, ox, Inches(1.65), pw, Inches(0.42), fill=hdr_color)
    txt(sl, f"  {title}", ox, Inches(1.65), pw, Inches(0.42),
        size=13, bold=True, color=WHITE)

    # 발신 (Outbound)
    rect(sl, ox, Inches(2.12), pw, Inches(0.36), fill=RGBColor(0x1a, 0x42, 0x2e))
    txt(sl, "  ▲  발신 (Outbound) — 내 서버  →  Knox 수신 서버",
        ox, Inches(2.12), pw, Inches(0.36), size=11, bold=True, color=STR_FG)
    out_lines = [
        ("출발지", "내 서버 IP   :   Port 80"),
        ("목적지", out_ip),
        ("Host",   out_host),
        ("포트",   "80,  443"),
    ]
    for j, (k, v) in enumerate(out_lines):
        by = Inches(2.52) + j * Inches(0.38)
        rect(sl, ox, by, Inches(1.3), Inches(0.38), fill=RGBColor(0x0c, 0x2a, 0x18))
        txt(sl, k, ox + Inches(0.1), by + Inches(0.04), Inches(1.2), Inches(0.32),
            size=11, bold=True, color=GREEN)
        txt(sl, v, ox + Inches(1.35), by + Inches(0.04), pw - Inches(1.4), Inches(0.32),
            size=11, color=WHITE)

    # 수신 (Inbound) — 빨간색 → 슬레이트
    ib_y = Inches(4.05)
    rect(sl, ox, ib_y, pw, Inches(0.36), fill=_SLATE)
    txt(sl, "  ▼  수신 (Inbound) — Knox 발신 서버  →  내 서버  ⚠ 반드시 허용",
        ox, ib_y, pw, Inches(0.36), size=11, bold=True, color=WHITE)
    in_lines = [
        ("출발지", in_ips),
        ("목적지", "내 서버 IP"),
        ("포트",   "80"),
    ]
    for j, (k, v) in enumerate(in_lines):
        by = ib_y + Inches(0.4) + j * Inches(0.42)
        rect(sl, ox, by, Inches(1.3), Inches(0.42), fill=RGBColor(0x2a, 0x32, 0x44))
        txt(sl, k, ox + Inches(0.1), by + Inches(0.05), Inches(1.2), Inches(0.34),
            size=11, bold=True, color=LBLUE)
        txt(sl, v, ox + Inches(1.35), by + Inches(0.05), pw - Inches(1.4), Inches(0.34),
            size=11, color=WHITE)

fw_panel(sl, Inches(0.4), "🔷  스테이지 (개발 서버)",
         MBLUE,
         "203.254.214.131",
         "openapi.stage.samsung.net",
         "112.106.197.162  (발신 서버 1개)")

fw_panel(sl, Inches(6.93), "🔷  운영 (실사용 서버)",
         BLUE,
         "112.107.220.134",
         "openapi.samsung.net",
         "182.195.35.14\n182.195.35.15\n182.195.35.16  (발신 서버 3개)")


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 6 — STEP 03: Knox 서버에 등록하기
# ══════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
header(sl, "STEP 03", "Knox 수/발신 서버에 등록하기",
       "봇을 Knox 서버에 등록하면 System ID와 Access Token이 발급됩니다 (본인만 봇 앱 리스트 확인 가능)")

# 좌: 스테이지
lx = Inches(0.4)
rect(sl, lx, Inches(1.65), Inches(6.1), Inches(0.42), fill=MBLUE)
txt(sl, "  🔷  스테이지 봇 연계 신청",
    lx, Inches(1.65), Inches(6.1), Inches(0.42), size=13, bold=True, color=WHITE)

path_box(sl, lx, Inches(2.12), Inches(6.1),
         "Home → My Interface → Bot → 연계 신청관리 → 스테이지 봇 연계 신청")

st_items = [
    ("결재",            "그룹장"),
    ("합의",            "없음"),
    ("AI빌더 사용여부", "미사용  (AI빌더 서버 주소가 달라 신중히 선택)"),
    ("과금 관리자",     "본인"),
]
for j, (k, v) in enumerate(st_items):
    by = Inches(2.72) + j * Inches(0.5)
    txt(sl, f"• {k}:", lx + Inches(0.2), by, Inches(1.9), Inches(0.45),
        size=12, bold=True, color=ACCENT)
    txt(sl, v, lx + Inches(2.15), by, Inches(4.2), Inches(0.45),
        size=12, color=WHITE)

photo_box(sl, lx, Inches(4.78), Inches(6.1), Inches(2.35),
          "스테이지 봇 연계 신청 화면 캡쳐")

# 우: 운영
rx = Inches(6.9)
rect(sl, rx, Inches(1.65), Inches(6.1), Inches(0.42),
     fill=RGBColor(0x5c, 0x35, 0x05))
txt(sl, "  🔶  운영 봇 연계 신청",
    rx, Inches(1.65), Inches(6.1), Inches(0.42), size=13, bold=True, color=WHITE)

path_box(sl, rx, Inches(2.12), Inches(6.1),
         "Home → My Interface → Bot → 연계 신청관리 → 운영 봇 연계 신청")

op_items = [
    ("AI빌더 사용여부", "미사용  (AI빌더 서버 주소가 달라 신중히 선택)"),
    ("과금 관리자",     "본인"),
]
for j, (k, v) in enumerate(op_items):
    by = Inches(2.72) + j * Inches(0.5)
    txt(sl, f"• {k}:", rx + Inches(0.2), by, Inches(1.9), Inches(0.45),
        size=12, bold=True, color=ORANGE)
    txt(sl, v, rx + Inches(2.15), by, Inches(4.2), Inches(0.45),
        size=12, color=WHITE)

note_box(sl, rx, Inches(3.8), Inches(6.1),
         "⚠  등록 완료 시 System ID 와 Access Token 을 꼭 보관하세요", color=ORANGE)

photo_box(sl, rx, Inches(4.42), Inches(6.1), Inches(2.7),
          "운영 봇 연계 신청 화면 캡쳐")


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 7 — STEP 04: API 연결 & 챗봇 방 생성
# ══════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
header(sl, "STEP 04", "API 연결 & 챗봇 방 생성",
       "Knox Developer Center의 API 문서를 참고하여 챗봇 대화방을 생성합니다")

path_box(sl, Inches(0.4), Inches(1.65), W - Inches(0.8),
         "Home  →  Dev Guide  →  메신저   또는   Home  →  Dev Guide  →  Bot  →  시작하기")

# 좌: 절차
lx = Inches(0.4)
steps_4 = [
    ("①", "메신저 Device 등록",
     "봇 계정에 Device를 등록합니다.\nSystem ID 와 Access Token 을 사용합니다."),
    ("②", "대화방 생성",
     "등록된 Device 로 챗봇 대화방을 생성합니다."),
    ("③", "API 문서 확인 및 연동",
     "Messenger API 문서를 참고하여\n발신/수신 API 를 개발 서버에 연결합니다."),
]
for j, (num, title_s, desc) in enumerate(steps_4):
    by = Inches(2.3) + j * Inches(1.45)
    rect(sl, lx, by, Inches(5.9), Inches(1.3), fill=BLUE)
    rect(sl, lx, by, Inches(5.9), Inches(0.08), fill=ACCENT)
    txt(sl, num, lx + Inches(0.1), by + Inches(0.1),
        Inches(0.6), Inches(0.55), size=22, bold=True, color=ACCENT)
    txt(sl, title_s, lx + Inches(0.75), by + Inches(0.1),
        Inches(5.0), Inches(0.45), size=14, bold=True, color=WHITE)
    txt(sl, desc, lx + Inches(0.75), by + Inches(0.55),
        Inches(5.0), Inches(0.7), size=11, color=LBLUE)

note_box(sl, lx, Inches(6.68), Inches(5.9),
         "📎 API 문서: Developer Center에서 Messenger API 문서를 첨부하세요", color=ACCENT)

# 우: 캡쳐 박스
photo_box(sl, Inches(6.7), Inches(2.3), Inches(6.3), Inches(4.9),
          "Dev Guide 화면 캡쳐")


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 7b — STEP 04: API 코드 예제 (Python)
# ══════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
header(sl, "STEP 04", "API 코드 예제 — Python",
       "① Device 등록  →  ② 대화방 생성  →  ③ 메시지 전송 순서로 호출합니다")

_lx = Inches(0.4)
_rx = Inches(6.93)
_cw = Inches(6.1)


def _cbox(sl, ox, oy, w, h, title, title_fill, code_str):
    rect(sl, ox, oy, w, Inches(0.32), fill=title_fill)
    txt(sl, f"  {title}", ox, oy, w, Inches(0.32), size=11, bold=True, color=WHITE)
    rect(sl, ox, oy + Inches(0.32), w, h, fill=CODE_BG)
    rect(sl, ox, oy + Inches(0.32), Inches(0.05), h, fill=ACCENT)
    txt(sl, code_str,
        ox + Inches(0.12), oy + Inches(0.37),
        w - Inches(0.18), h - Inches(0.08),
        size=9.5, color=CODE_FG)


# ── 공통 설정 ──────────────────────────────────────────────────────────────
_c0 = (
    'import requests\n\n'
    'BASE = "https://openapi.stage.samsung.net"\n'
    'headers = {\n'
    '    "Authorization": "Bearer <ACCESS_TOKEN>",\n'
    '    "x-system-id":   "<SYSTEM_ID>",\n'
    '    "Content-Type":  "application/json",\n'
    '}'
)
_cbox(sl, _lx, Inches(1.65), _cw, Inches(1.65), "공통 설정 (Base URL & Headers)", MBLUE, _c0)

# ── ① Device 등록 ──────────────────────────────────────────────────────────
_c1 = (
    'res = requests.post(\n'
    '    f"{BASE}/messenger/bot/v1/device",\n'
    '    headers=headers,\n'
    '    json={"deviceName": "fa-service-bot"},\n'
    ')\n'
    'device_id = res.json()["deviceId"]\n'
    '# device_id 를 저장해 두세요'
)
_cbox(sl, _lx, Inches(3.72), _cw, Inches(1.65), "① Device 등록", BLUE, _c1)

note_box(sl, _lx, Inches(5.52), _cw,
         "💡 device_id : 이후 모든 API 호출에  x-device-id  헤더로 사용합니다", color=GREEN)
note_box(sl, _lx, Inches(6.12), _cw,
         "📌 운영 서버:  BASE = \"https://openapi.samsung.net\"  으로 변경", color=DGRAY)

# ── ② 대화방 생성 ──────────────────────────────────────────────────────────
_c2 = (
    'res = requests.post(\n'
    '    f"{BASE}/messenger/bot/v1/chatroom",\n'
    '    headers={**headers, "x-device-id": device_id},\n'
    '    json={\n'
    '        "roomName": "FA 분석 알림",\n'
    '        "memberList": ["hong.gildong@samsung.com"],\n'
    '    },\n'
    ')\n'
    'room_id = res.json()["roomId"]'
)
_cbox(sl, _rx, Inches(1.65), _cw, Inches(2.05), "② 대화방 생성", BLUE, _c2)

# ── ③ 메시지 전송 ──────────────────────────────────────────────────────────
_c3 = (
    'res = requests.post(\n'
    '    f"{BASE}/messenger/bot/v1/message",\n'
    '    headers={**headers, "x-device-id": device_id},\n'
    '    json={\n'
    '        "roomId": room_id,\n'
    '        "messageType": "TEXT",\n'
    '        "message": "FA 분석이 완료되었습니다.",\n'
    '    },\n'
    ')\n'
    '# 응답: {"result": "success", ...}'
)
_cbox(sl, _rx, Inches(3.82), _cw, Inches(2.05), "③ 메시지 전송", MBLUE, _c3)

note_box(sl, _rx, Inches(6.00), _cw,
         "💡 실제 엔드포인트/파라미터는 Knox Developer Center 공식 문서를 반드시 확인하세요",
         color=ACCENT)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 7c — STEP 04: Adaptive Card & PDF 파일 전송 예제
# ══════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
header(sl, "STEP 04", "Adaptive Card & 파일 전송 예제",
       "카드 메시지로 링크 버튼을 포함한 알림을 보내고, PDF 파일도 직접 전송할 수 있습니다")

_lx2 = Inches(0.4)
_rx2 = Inches(6.93)
_cw2 = Inches(6.1)

# ── Adaptive Card 발신 ─────────────────────────────────────────────────────
_card = (
    'res = requests.post(\n'
    '    f"{BASE}/messenger/bot/v1/message",\n'
    '    headers={**headers, "x-device-id": device_id},\n'
    '    json={\n'
    '        "roomId": room_id,\n'
    '        "messageType": "CARD",\n'
    '        "cardTitle": "FA 분석 결과",\n'
    '        "cardContent": "SN: R3CX1234...\\n이상 항목: 3건",\n'
    '        "cardButtonList": [\n'
    '            {\n'
    '                "buttonType": "LINK",\n'
    '                "buttonText": "결과 보기",\n'
    '                "buttonValue": "http://10.246.9.74/result",\n'
    '            },\n'
    '        ],\n'
    '    },\n'
    ')'
)
rect(sl, _lx2, Inches(1.65), _cw2, Inches(0.32), fill=RGBColor(0x14, 0x53, 0x6e))
txt(sl, "  📋  Adaptive Card 발신 (링크 버튼 포함)",
    _lx2, Inches(1.65), _cw2, Inches(0.32), size=11, bold=True, color=WHITE)
rect(sl, _lx2, Inches(1.97), _cw2, Inches(3.25), fill=CODE_BG)
rect(sl, _lx2, Inches(1.97), Inches(0.05), Inches(3.25), fill=ACCENT)
txt(sl, _card,
    _lx2 + Inches(0.12), Inches(2.02),
    _cw2 - Inches(0.18), Inches(3.18),
    size=9.5, color=CODE_FG)

note_box(sl, _lx2, Inches(5.35), _cw2,
         "💡 cardButtonList 에 버튼 여러 개 추가 가능 (LINK / 전화 / 위치 등)", color=LBLUE)
note_box(sl, _lx2, Inches(5.97), _cw2,
         "📌 cardContent 줄바꿈:  \\n  사용 (실제 문자열에서 \\\\n → \\n)", color=DGRAY)

# ── PDF 파일 전송 ──────────────────────────────────────────────────────────
_pdf = (
    '# ① 파일 업로드 → fileId 획득\n'
    'with open("fa_report.pdf", "rb") as f:\n'
    '    res = requests.post(\n'
    '        f"{BASE}/messenger/bot/v1/file",\n'
    '        headers={**headers, "x-device-id": device_id},\n'
    '        files={"file": ("fa_report.pdf", f, "application/pdf")},\n'
    '    )\n'
    'file_id = res.json()["fileId"]\n'
    '\n'
    '# ② 파일 메시지 전송\n'
    'res = requests.post(\n'
    '    f"{BASE}/messenger/bot/v1/message",\n'
    '    headers={**headers, "x-device-id": device_id},\n'
    '    json={\n'
    '        "roomId": room_id,\n'
    '        "messageType": "FILE",\n'
    '        "fileId": file_id,\n'
    '    },\n'
    ')'
)
rect(sl, _rx2, Inches(1.65), _cw2, Inches(0.32), fill=RGBColor(0x4a, 0x27, 0x6e))
txt(sl, "  📎  PDF 파일 전송 (업로드 → 메시지)",
    _rx2, Inches(1.65), _cw2, Inches(0.32), size=11, bold=True, color=WHITE)
rect(sl, _rx2, Inches(1.97), _cw2, Inches(3.85), fill=CODE_BG)
rect(sl, _rx2, Inches(1.97), Inches(0.05), Inches(3.85), fill=RGBColor(0x9b, 0x5d, 0xe5))
txt(sl, _pdf,
    _rx2 + Inches(0.12), Inches(2.02),
    _cw2 - Inches(0.18), Inches(3.78),
    size=9.5, color=CODE_FG)

note_box(sl, _rx2, Inches(5.95), _cw2,
         "💡 파일 크기 제한 및 허용 확장자는 Knox Developer Center 문서를 확인하세요", color=RGBColor(0x9b, 0x5d, 0xe5))


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 8 — STEP 05: FAQ
# ══════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
header(sl, "STEP 05", "FAQ",
       "궁금한 점이 있을 때 참고하세요")

faq = [
    ("Q1", "문의는 어디에 하나요?",
     "Knoxportal@samsung.com 으로 이메일 문의"),
    ("Q2", "Knox Support는 어떻게 사용하나요?",
     "Knox Portal 화면 내 프로필 왼편  ❓  → Knox Support 진입\n"
     "Teams 메뉴에서 검색 및 문의 업로드 가능"),
]

for j, (num, q, a) in enumerate(faq):
    by = Inches(1.75) + j * Inches(1.8)
    bw2 = Inches(6.2)
    rect(sl, Inches(0.4), by, bw2, Inches(1.55), fill=BLUE)
    rect(sl, Inches(0.4), by, bw2, Inches(0.08), fill=ACCENT)
    rect(sl, Inches(0.4), by + Inches(0.12), Inches(0.6), Inches(0.55),
         fill=MBLUE)
    txt(sl, num, Inches(0.4), by + Inches(0.12), Inches(0.6), Inches(0.52),
        size=13, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
    txt(sl, q, Inches(1.1), by + Inches(0.12), bw2 - Inches(0.75), Inches(0.52),
        size=14, bold=True, color=WHITE)
    txt(sl, a, Inches(1.1), by + Inches(0.7), bw2 - Inches(0.75), Inches(0.75),
        size=12, color=LBLUE)

photo_box(sl, Inches(7.1), Inches(1.75), Inches(5.9), Inches(5.3),
          "Knox Support 화면 캡쳐")

txt(sl, "Knoxportal@samsung.com",
    Inches(0.5), Inches(5.5), Inches(6.0), Inches(0.5),
    size=14, bold=True, color=ACCENT)

note_box(sl, Inches(0.4), Inches(6.1), Inches(6.5),
         "💡 Knox Portal 내 물음표(?) 버튼 → Knox Support → Teams 메뉴", color=LBLUE)


# ══════════════════════════════════════════════════════════════════════════════
# 저장
# ══════════════════════════════════════════════════════════════════════════════
out = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "knox_chatbot_가이드.pptx"
)
prs.save(out)
print(f"저장 완료: {out}")
