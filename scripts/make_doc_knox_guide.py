"""
Knox Teams 메신저 만들기 - 가이드 문서 생성 (python-docx)
"""
import os
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

for sec in doc.sections:
    sec.top_margin    = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    sec.left_margin   = Cm(3.0)
    sec.right_margin  = Cm(2.5)

# ── 색상 ─────────────────────────────────────────────────────────────────────
NAVY    = (0x1e, 0x3a, 0x6e)
BLUE    = (0x2d, 0x5a, 0x9e)
WHITE   = (0xFF, 0xFF, 0xFF)
GREEN_D = (0x1e, 0x88, 0x55)
RED_D   = (0xc0, 0x39, 0x2b)

# ── 헬퍼 ─────────────────────────────────────────────────────────────────────
def _shd(cell, hex6):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn('w:shd')):
        tcPr.remove(old)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex6)
    tcPr.append(shd)


def _cp(cell, text, bold=False, size=10, rgb=None,
        align=WD_ALIGN_PARAGRAPH.LEFT, italic=False):
    p = cell.paragraphs[0]
    p.alignment = align
    run = p.add_run(text)
    run.font.bold   = bold
    run.font.italic = italic
    run.font.size   = Pt(size)
    if rgb:
        run.font.color.rgb = RGBColor(*rgb)


def heading(level, text, rgb=NAVY):
    p = doc.add_heading(text, level=level)
    if p.runs:
        p.runs[0].font.color.rgb = RGBColor(*rgb)
    return p


def screenshot_box(label="캡쳐 필요"):
    tbl  = doc.add_table(rows=1, cols=1)
    tbl.style = 'Table Grid'
    cell = tbl.cell(0, 0)
    _shd(cell, 'EBF5FB')
    tr   = cell._tc.getparent()
    trPr = tr.get_or_add_trPr()
    trH  = OxmlElement('w:trHeight')
    trH.set(qn('w:val'), '1400')
    trH.set(qn('w:hRule'), 'exact')
    trPr.append(trH)
    p    = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run  = p.add_run(f'[ 📷  {label} ]')
    run.font.size  = Pt(11)
    run.font.bold  = True
    run.font.color.rgb = RGBColor(0x2e, 0x86, 0xc1)
    doc.add_paragraph()


def path_box(text):
    tbl  = doc.add_table(rows=1, cols=1)
    tbl.style = 'Table Grid'
    cell = tbl.cell(0, 0)
    _shd(cell, 'EAF2FF')
    p    = cell.paragraphs[0]
    p.paragraph_format.left_indent = Cm(0.3)
    run  = p.add_run(text)
    run.font.bold  = True
    run.font.size  = Pt(10)
    run.font.color.rgb = RGBColor(*NAVY)
    doc.add_paragraph()


def note_box(text, bg='FEF9E7', rgb=(0x78, 0x50, 0x10)):
    tbl  = doc.add_table(rows=1, cols=1)
    tbl.style = 'Table Grid'
    cell = tbl.cell(0, 0)
    _shd(cell, bg)
    p    = cell.paragraphs[0]
    p.paragraph_format.left_indent = Cm(0.3)
    run  = p.add_run(text)
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(*rgb)
    doc.add_paragraph()


def bullet(text='', level=0, bold_part=None, extra=''):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent  = Cm(0.8 + level * 0.8)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(2)
    run0 = p.add_run('•  ')
    run0.font.color.rgb = RGBColor(*BLUE)
    run0.font.size = Pt(10)
    if bold_part:
        r1 = p.add_run(bold_part)
        r1.font.bold = True
        r1.font.size = Pt(10)
        r2 = p.add_run(extra or text)
        r2.font.size = Pt(10)
    else:
        r = p.add_run(text)
        r.font.size = Pt(10)


# ══════════════════════════════════════════════════════════════════════════════
# 제목 페이지
# ══════════════════════════════════════════════════════════════════════════════
t = doc.add_heading('Knox Teams 메신저 만들기', level=0)
if t.runs:
    t.runs[0].font.color.rgb = RGBColor(*NAVY)
t.alignment = WD_ALIGN_PARAGRAPH.CENTER

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run('Knox C&C Suite Developer Guide')
r.font.size = Pt(13)
r.font.color.rgb = RGBColor(*BLUE)
r.font.italic = True

doc.add_paragraph()

dev = doc.add_paragraph()
r1 = dev.add_run('개발 사이트: Knox C&C suite Developers Center\n')
r1.font.bold = True
r1.font.size = Pt(10)
r2 = dev.add_run('http://developers.samsung.net/static/knoxcenter/main.html#')
r2.font.size = Pt(10)
r2.font.color.rgb = RGBColor(0x27, 0x6d, 0xc6)

doc.add_page_break()


# ══════════════════════════════════════════════════════════════════════════════
# 1. 챗봇 아이디 생성하기
# ══════════════════════════════════════════════════════════════════════════════
heading(1, '1. 챗봇 아이디 생성하기')

doc.add_paragraph('Knox Portal에서 봇 계정을 신청합니다. 결재 후 챗봇 아이디가 생성됩니다.')
doc.add_paragraph()

path_box('Home  →  My Interface  →  Bot  →  봇관리  →  봇계정신청')

p = doc.add_paragraph()
r = p.add_run('결재 정보')
r.font.bold = True
r.font.size = Pt(10)
bullet('파트장 결재')
bullet('경영지원 합의')

doc.add_paragraph()
screenshot_box('봇계정신청 화면 캡쳐')

doc.add_page_break()


# ══════════════════════════════════════════════════════════════════════════════
# 2. 방화벽 연결하기
# ══════════════════════════════════════════════════════════════════════════════
heading(1, '2. 방화벽 연결하기')

doc.add_paragraph(
    '내 서버와 Knox Messenger 서버 간 통신이 가능하도록 방화벽 정책을 등록합니다.\n'
    '발신(Outbound)과 수신(Inbound) 두 방향 모두 설정이 필요합니다.'
)
doc.add_paragraph()

# ── 개념도 ────────────────────────────────────────────────────────────────────
heading(3, '▌ 방화벽 개념도', rgb=BLUE)

COL_W = [3.0, 1.5, 2.5, 1.5, 3.0]  # cm

diag = doc.add_table(rows=4, cols=5)
diag.style = 'Table Grid'
for ri, row in enumerate(diag.rows):
    for ci, w in enumerate(COL_W):
        row.cells[ci].width = Cm(w)

# Row 0 — 헤더
R = diag.rows[0].cells
for c in R: _shd(c, '1E3A6E')
_cp(R[0], '내 서버',    bold=True, size=11, rgb=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
_cp(R[1], '',           size=10,   rgb=WHITE)
_cp(R[2], '🔥 방화벽', bold=True, size=11, rgb=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
_cp(R[3], '',           size=10,   rgb=WHITE)
_cp(R[4], 'Knox 서버',  bold=True, size=11, rgb=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)

# Row 1 — 발신 (Outbound)
R = diag.rows[1].cells
for c in R: _shd(c, 'E8F8F5')
_cp(R[0], '내 서버 IP\nPort 80',           size=9,  align=WD_ALIGN_PARAGRAPH.CENTER)
_cp(R[1], '────▶',  bold=True, size=13, rgb=GREEN_D, align=WD_ALIGN_PARAGRAPH.CENTER)
_cp(R[2], '발신 (Outbound)\nPort 80 / 443', size=9, rgb=GREEN_D, align=WD_ALIGN_PARAGRAPH.CENTER)
_cp(R[3], '────▶',  bold=True, size=13, rgb=GREEN_D, align=WD_ALIGN_PARAGRAPH.CENTER)
_cp(R[4], 'openapi.samsung.net\n(or stage)', size=9, align=WD_ALIGN_PARAGRAPH.CENTER)

# Row 2 — 수신 (Inbound)
R = diag.rows[2].cells
for c in R: _shd(c, 'FEF5E7')
_cp(R[0], '내 서버 IP\nPort 80',             size=9,  align=WD_ALIGN_PARAGRAPH.CENTER)
_cp(R[1], '◀────',  bold=True, size=13, rgb=RED_D, align=WD_ALIGN_PARAGRAPH.CENTER)
_cp(R[2], '수신 (Inbound)\n반드시 허용!', size=9, rgb=RED_D, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
_cp(R[3], '◀────',  bold=True, size=13, rgb=RED_D, align=WD_ALIGN_PARAGRAPH.CENTER)
_cp(R[4], 'Knox 서버 IP\n(출발지)',           size=9,  align=WD_ALIGN_PARAGRAPH.CENTER)

# Row 3 — 노트
R = diag.rows[3].cells
merged = R[0].merge(R[4])
_shd(merged, 'FDFEFE')
_cp(merged,
    '⚠  수신(Inbound) 정책: Knox 서버 IP에서 내 서버 Port 80으로 들어오는 트래픽을 허용해야 메시지를 받을 수 있습니다.',
    size=9, rgb=(0x78, 0x28, 0x28))

doc.add_paragraph()

# ── 2-1 스테이지 ──────────────────────────────────────────────────────────────
heading(2, '2-1. 스테이지 (개발 서버)', rgb=BLUE)

fw_tbl = doc.add_table(rows=3, cols=5)
fw_tbl.style = 'Table Grid'
HDR = ['방향', '출발지 (From)', '', '목적지 (To)', '포트']
HDR_W = [1.8, 3.5, 0.8, 4.0, 1.2]

for ri, row in enumerate(fw_tbl.rows):
    for ci, w in enumerate(HDR_W):
        row.cells[ci].width = Cm(w)

# header
R = fw_tbl.rows[0].cells
for ci, h in enumerate(HDR):
    _shd(R[ci], '2D5A9E')
    _cp(R[ci], h, bold=True, size=10, rgb=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)

# outbound
R = fw_tbl.rows[1].cells
for c in R: _shd(c, 'F0FFF4')
_cp(R[0], '발신\n(Outbound)', bold=True, size=10, rgb=GREEN_D, align=WD_ALIGN_PARAGRAPH.CENTER)
_cp(R[1], '내 서버 IP  :  Port 80', size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
_cp(R[2], '→', bold=True, size=14, rgb=GREEN_D, align=WD_ALIGN_PARAGRAPH.CENTER)
_cp(R[3], '203.254.214.131\nopenapi.stage.samsung.net', size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
_cp(R[4], '80, 443', size=10, rgb=GREEN_D, align=WD_ALIGN_PARAGRAPH.CENTER)

# inbound
R = fw_tbl.rows[2].cells
for c in R: _shd(c, 'FFF9F0')
_cp(R[0], '수신\n(Inbound)', bold=True, size=10, rgb=RED_D, align=WD_ALIGN_PARAGRAPH.CENTER)
_cp(R[1], '112.106.197.162\n(Knox 서버 IP, 1개)', size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
_cp(R[2], '→', bold=True, size=14, rgb=RED_D, align=WD_ALIGN_PARAGRAPH.CENTER)
_cp(R[3], '내 서버 IP', size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
_cp(R[4], '80', size=10, rgb=RED_D, align=WD_ALIGN_PARAGRAPH.CENTER)

doc.add_paragraph()

# ── 2-2 운영 ──────────────────────────────────────────────────────────────────
heading(2, '2-2. 운영 (실사용 서버)', rgb=BLUE)

fw_tbl2 = doc.add_table(rows=3, cols=5)
fw_tbl2.style = 'Table Grid'
for ri, row in enumerate(fw_tbl2.rows):
    for ci, w in enumerate(HDR_W):
        row.cells[ci].width = Cm(w)

R = fw_tbl2.rows[0].cells
for ci, h in enumerate(HDR):
    _shd(R[ci], '1E3A6E')
    _cp(R[ci], h, bold=True, size=10, rgb=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)

R = fw_tbl2.rows[1].cells
for c in R: _shd(c, 'F0FFF4')
_cp(R[0], '발신\n(Outbound)', bold=True, size=10, rgb=GREEN_D, align=WD_ALIGN_PARAGRAPH.CENTER)
_cp(R[1], '내 서버 IP  :  Port 80', size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
_cp(R[2], '→', bold=True, size=14, rgb=GREEN_D, align=WD_ALIGN_PARAGRAPH.CENTER)
_cp(R[3], '112.107.220.134\nopenapi.samsung.net', size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
_cp(R[4], '80, 443', size=10, rgb=GREEN_D, align=WD_ALIGN_PARAGRAPH.CENTER)

R = fw_tbl2.rows[2].cells
for c in R: _shd(c, 'FFF9F0')
_cp(R[0], '수신\n(Inbound)', bold=True, size=10, rgb=RED_D, align=WD_ALIGN_PARAGRAPH.CENTER)
p_ip = R[1].paragraphs[0]
p_ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
rr = p_ip.add_run('182.195.35.14\n182.195.35.15\n182.195.35.16\n(Knox 서버 IP, 3개)')
rr.font.size = Pt(9)
_cp(R[2], '→', bold=True, size=14, rgb=RED_D, align=WD_ALIGN_PARAGRAPH.CENTER)
_cp(R[3], '내 서버 IP', size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
_cp(R[4], '80', size=10, rgb=RED_D, align=WD_ALIGN_PARAGRAPH.CENTER)

doc.add_paragraph()
note_box('💡 방화벽 정책 등록 후 담당 인프라팀에 확인을 받으세요.')

doc.add_page_break()


# ══════════════════════════════════════════════════════════════════════════════
# 3. Knox 수/발신 서버에 등록하기
# ══════════════════════════════════════════════════════════════════════════════
heading(1, '3. Knox 수/발신 서버에 등록하기')

doc.add_paragraph('봇을 Knox 서버에 등록하면 System ID와 Access Token이 발급됩니다.\n'
                   '(등록 후 본인 계정에서만 챗봇 앱 리스트에서 확인 가능)')
doc.add_paragraph()

# 스테이지
heading(2, '스테이지 봇 연계 신청', rgb=BLUE)
path_box('Home  →  My Interface  →  Bot  →  연계 신청관리  →  스테이지 봇 연계 신청')

p = doc.add_paragraph()
r = p.add_run('결재 정보')
r.font.bold = True; r.font.size = Pt(10)
bullet('결재: 그룹장')
bullet('합의: 없음')

doc.add_paragraph()
p = doc.add_paragraph()
r = p.add_run('신청 시 주의 사항')
r.font.bold = True; r.font.size = Pt(10)
bullet(bold_part='미사용', extra=' — AI빌더 사용여부: 미사용 (AI빌더 관리 서버와 일반 개발 서버 주소가 다르므로 신중히 선택)')
bullet(bold_part='본인', extra=' — 과금 관리자: 본인')

doc.add_paragraph()
screenshot_box('스테이지 봇 연계 신청 화면 캡쳐')

# 운영
heading(2, '운영 봇 연계 신청', rgb=BLUE)

p = doc.add_paragraph()
r = p.add_run('신청 시 주의 사항')
r.font.bold = True; r.font.size = Pt(10)
bullet(bold_part='미사용', extra=' — AI빌더 사용여부: 미사용 (AI빌더 관리 서버와 일반 개발 서버 주소가 다르므로 신중히 선택)')
bullet(bold_part='본인', extra=' — 과금 관리자: 본인')

doc.add_paragraph()
screenshot_box('운영 봇 연계 신청 화면 캡쳐')

doc.add_page_break()


# ══════════════════════════════════════════════════════════════════════════════
# 4. API 연결해서 챗봇 방 생성하기
# ══════════════════════════════════════════════════════════════════════════════
heading(1, '4. API 연결해서 챗봇 방 생성하기')

path_box('Home  →  Dev Guide  →  메신저   또는   Home  →  Dev Guide  →  Bot  →  시작하기')

screenshot_box('Dev Guide 화면 캡쳐')

heading(2, '챗봇 대화방 생성 예시', rgb=BLUE)

steps = [
    ('① 메신저 Device 등록',
     '봇 계정에 Device를 등록합니다. System ID와 Access Token을 사용합니다.'),
    ('② 대화방 생성',
     '등록된 Device로 챗봇 대화방을 생성합니다.'),
    ('③ API 문서 확인 및 연동',
     'Knox Developer Center의 Messenger API 문서를 참고하여 발신/수신 API를 연결합니다.'),
]
for title_s, desc in steps:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    r1 = p.add_run(title_s + '  ')
    r1.font.bold = True
    r1.font.size = Pt(11)
    r1.font.color.rgb = RGBColor(*BLUE)
    r2 = p.add_run(desc)
    r2.font.size = Pt(10)

doc.add_paragraph()
note_box('📎 API 문서: Developer Center에서 Messenger API 문서를 다운로드하여 첨부하세요.',
         bg='EBF5FB', rgb=(0x1a, 0x53, 0x8b))

doc.add_page_break()


# ══════════════════════════════════════════════════════════════════════════════
# 5. FAQ
# ══════════════════════════════════════════════════════════════════════════════
heading(1, '5. FAQ')

faq = [
    ('Q. 문의는 어디에 하나요?',
     'Knoxportal@samsung.com 으로 이메일 문의'),
    ('Q. Knox Support는 어떻게 사용하나요?',
     'Knox Portal 화면 내 프로필 왼편 물음표(?) → Knox Support 진입\n'
     'Teams 메뉴에서 검색 및 업로드 가능'),
]
for q, a in faq:
    p = doc.add_paragraph()
    r = p.add_run(q)
    r.font.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(*NAVY)

    p2 = doc.add_paragraph()
    p2.paragraph_format.left_indent = Cm(1.0)
    r2 = p2.add_run(a)
    r2.font.size = Pt(10)
    doc.add_paragraph()

screenshot_box('Knox Support 화면 캡쳐')


# ── 저장 ─────────────────────────────────────────────────────────────────────
out = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "knox_chatbot_가이드.docx"
)
doc.save(out)
print(f"저장 완료: {out}")
