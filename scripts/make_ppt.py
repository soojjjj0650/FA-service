"""
Knox 챗봇 서비스 소개 PPT 생성기
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
import pptx.oxml.ns as nsmap
from lxml import etree
import copy

# ── 색상 정의 ──────────────────────────────────────────────────────────────────
NAVY    = RGBColor(0x1e, 0x3a, 0x5f)
BLUE    = RGBColor(0x2d, 0x5a, 0x9e)
LBLUE   = RGBColor(0x60, 0xa5, 0xfa)
ACCENT  = RGBColor(0x38, 0xbd, 0xf8)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
GRAY    = RGBColor(0xf1, 0xf5, 0xf9)
DGRAY   = RGBColor(0x64, 0x74, 0x8b)
BLACK   = RGBColor(0x0f, 0x17, 0x2a)
GREEN   = RGBColor(0x10, 0xb9, 0x81)
ORANGE  = RGBColor(0xf5, 0x9e, 0x0b)

W = Inches(13.33)
H = Inches(7.5)

prs = Presentation()
prs.slide_width  = W
prs.slide_height = H

BLANK = prs.slide_layouts[6]  # completely blank


def add_rect(slide, x, y, w, h, fill=None, line=None, alpha=None):
    shape = slide.shapes.add_shape(1, x, y, w, h)
    shape.line.fill.background()
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def add_text(slide, text, x, y, w, h,
             size=24, bold=False, color=BLACK, align=PP_ALIGN.LEFT,
             italic=False, wrap=True):
    txb = slide.shapes.add_textbox(x, y, w, h)
    tf  = txb.text_frame
    tf.word_wrap = wrap
    p   = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size  = Pt(size)
    run.font.bold  = bold
    run.font.color.rgb = color
    run.font.italic = italic
    return txb


def add_para(tf, text, size=18, bold=False, color=BLACK,
             align=PP_ALIGN.LEFT, space_before=0, italic=False):
    p = tf.add_paragraph()
    p.alignment = align
    if space_before:
        p.space_before = Pt(space_before)
    run = p.add_run()
    run.text = text
    run.font.size  = Pt(size)
    run.font.bold  = bold
    run.font.color.rgb = color
    run.font.italic = italic
    return p


def slide_header(slide, title, subtitle=None):
    """상단 헤더 바 + 제목"""
    add_rect(slide, 0, 0, W, Inches(0.12), fill=ACCENT)
    add_text(slide, title,
             Inches(0.55), Inches(0.22), Inches(11), Inches(0.7),
             size=28, bold=True, color=NAVY)
    if subtitle:
        add_text(slide, subtitle,
                 Inches(0.55), Inches(0.85), Inches(11), Inches(0.4),
                 size=15, color=DGRAY)
    add_rect(slide, Inches(0.4), Inches(1.15), W - Inches(0.8), Pt(1.5), fill=LBLUE)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 1 — 표지
# ══════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
add_rect(sl, 0, 0, W, H, fill=NAVY)
# 왼쪽 밝은 사이드바
add_rect(sl, 0, 0, Inches(0.25), H, fill=ACCENT)
# 배경 도형 장식
add_rect(sl, Inches(8.5), Inches(1.5), Inches(5.5), Inches(5.5), fill=BLUE)
add_rect(sl, Inches(9.5), Inches(2.2), Inches(4.5), Inches(4.5), fill=RGBColor(0x1a, 0x50, 0x8b))

add_text(sl, "📡", Inches(0.7), Inches(1.6), Inches(2), Inches(1.2), size=54, color=WHITE)

add_text(sl, "통화품질 분석\nFA 대응 챗봇 서비스",
         Inches(0.7), Inches(2.6), Inches(8), Inches(2.2),
         size=42, bold=True, color=WHITE)

add_text(sl, "Knox Messenger 기반 자동 분석 서비스 소개",
         Inches(0.7), Inches(4.8), Inches(8), Inches(0.6),
         size=18, color=LBLUE, italic=True)

add_text(sl, "무선 사업부  |  통화품질 개선 TF",
         Inches(0.7), Inches(6.5), Inches(6), Inches(0.5),
         size=13, color=RGBColor(0x94, 0xa3, 0xb8))


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 2 — 목차
# ══════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
add_rect(sl, 0, 0, W, H, fill=GRAY)
slide_header(sl, "목  차")

items = [
    ("01", "배경 및 문제",        "FA 미결건 처리의 어려움"),
    ("02", "서비스 개요",          "챗봇이 하는 일"),
    ("03", "주요 기능",            "어떤 기능이 있나요?"),
    ("04", "동작 흐름",            "SN 입력부터 결과까지"),
    ("05", "자동화 구성",          "매일 알아서 돌아가는 방식"),
    ("06", "기대 효과",            "도입 후 달라지는 것"),
]

cols = 2
col_w = Inches(5.8)
for i, (num, title, desc) in enumerate(items):
    col = i % cols
    row = i // cols
    bx = Inches(0.5) + col * Inches(6.7)
    by = Inches(1.5) + row * Inches(1.65)

    add_rect(sl, bx, by, col_w, Inches(1.4), fill=WHITE)
    add_rect(sl, bx, by, Inches(0.12), Inches(1.4), fill=ACCENT)

    add_text(sl, num, bx + Inches(0.25), by + Inches(0.12),
             Inches(0.7), Inches(0.5), size=22, bold=True, color=ACCENT)
    add_text(sl, title, bx + Inches(0.25), by + Inches(0.55),
             col_w - Inches(0.4), Inches(0.45), size=17, bold=True, color=NAVY)
    add_text(sl, desc, bx + Inches(0.25), by + Inches(0.95),
             col_w - Inches(0.4), Inches(0.35), size=12, color=DGRAY)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 3 — 배경 및 문제
# ══════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
add_rect(sl, 0, 0, W, H, fill=GRAY)
slide_header(sl, "배경 및 문제", "FA 미결건 처리, 이런 불편함이 있었습니다")

problems = [
    ("📧", "매일 아침\n이메일 수동 확인", "FA 미결건 엑셀 파일을\n매일 직접 다운로드"),
    ("🔍", "SN별 수동\n데이터 조회",        "시스템에서 하나씩\n직접 검색"),
    ("📊", "수작업\n분석 리포트",             "통화품질 데이터를\n수동으로 정리"),
    ("📨", "결과 전달\n지연",                 "분석 완료까지\n시간이 오래 소요"),
]

for i, (icon, title, desc) in enumerate(problems):
    bx = Inches(0.4) + i * Inches(3.18)
    by = Inches(1.5)

    add_rect(sl, bx, by, Inches(3.0), Inches(4.5), fill=WHITE)
    add_rect(sl, bx, by, Inches(3.0), Inches(0.08), fill=ORANGE)

    add_text(sl, icon, bx + Inches(0.9), by + Inches(0.3),
             Inches(1.2), Inches(1.0), size=40, align=PP_ALIGN.CENTER)
    add_text(sl, title, bx + Inches(0.1), by + Inches(1.35),
             Inches(2.8), Inches(0.9), size=15, bold=True, color=NAVY,
             align=PP_ALIGN.CENTER)
    add_text(sl, desc, bx + Inches(0.1), by + Inches(2.35),
             Inches(2.8), Inches(1.8), size=13, color=DGRAY,
             align=PP_ALIGN.CENTER)

add_text(sl, "💡  이 모든 과정을 자동화할 수 없을까?",
         Inches(1.5), Inches(6.4), Inches(10), Inches(0.6),
         size=16, bold=True, color=BLUE, align=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 4 — 서비스 개요
# ══════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
add_rect(sl, 0, 0, W, H, fill=GRAY)
slide_header(sl, "서비스 개요", "Knox 챗봇에 SN을 입력하면 분석 결과를 자동으로 받아볼 수 있습니다")

# 중앙 흐름도
steps = [
    ("💬", "SN 입력",     "Knox 챗봇에\n단말 SN 입력"),
    ("⚙️", "자동 분석",   "서버에서\n통화품질 데이터 조회"),
    ("📄", "PDF 생성",    "분석 결과\nPDF 자동 생성"),
    ("📱", "결과 전송",   "Knox 메신저로\n결과 PDF 수신"),
]

arrow = "  →  "
bw = Inches(2.6)
bh = Inches(3.5)
gap = Inches(0.55)
total = len(steps) * bw + (len(steps) - 1) * gap
sx = (W - total) / 2

for i, (icon, title, desc) in enumerate(steps):
    bx = sx + i * (bw + gap)
    by = Inches(1.7)

    add_rect(sl, bx, by, bw, bh, fill=NAVY)
    add_rect(sl, bx, by, bw, Inches(0.08), fill=ACCENT)

    add_text(sl, icon, bx, by + Inches(0.3),
             bw, Inches(0.9), size=36, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(sl, f"STEP {i+1}", bx, by + Inches(1.2),
             bw, Inches(0.4), size=11, color=ACCENT,
             align=PP_ALIGN.CENTER, bold=True)
    add_text(sl, title, bx, by + Inches(1.6),
             bw, Inches(0.6), size=16, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER)
    add_text(sl, desc, bx, by + Inches(2.2),
             bw, Inches(1.0), size=12, color=LBLUE,
             align=PP_ALIGN.CENTER)

    if i < len(steps) - 1:
        add_text(sl, "▶", bx + bw, by + Inches(1.5),
                 gap, Inches(0.6), size=22, color=ACCENT, align=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 5 — 주요 기능
# ══════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
add_rect(sl, 0, 0, W, H, fill=GRAY)
slide_header(sl, "주요 기능", "사용자 편의를 위한 다양한 기능을 제공합니다")

features = [
    ("🎯", "Adaptive Card UI",  "버튼/입력창 형태의\n직관적인 대화 인터페이스"),
    ("👥", "그룹채팅 지원",       "@봇 이름으로 호출\n여러 명이 함께 사용 가능"),
    ("📊", "통화품질 분석",       "MUTE·DROP·기지국별\n상세 분석 리포트"),
    ("📅", "일별/추이 분석",      "날짜별 상세 데이터\n트렌드 그래프 제공"),
    ("🗺️", "지도 시각화",         "문제 기지국 위치를\n지도에서 직접 확인"),
    ("⚡", "자동 스케줄링",       "매일 18:00 자동 실행\n별도 조작 불필요"),
]

cols = 3
for i, (icon, title, desc) in enumerate(features):
    col = i % cols
    row = i // cols
    fw = Inches(3.9)
    fh = Inches(2.2)
    bx = Inches(0.4) + col * Inches(4.18)
    by = Inches(1.5) + row * Inches(2.4)

    add_rect(sl, bx, by, fw, fh, fill=WHITE)
    add_rect(sl, bx, by, Inches(0.1), fh, fill=ACCENT)

    add_text(sl, icon, bx + Inches(0.25), by + Inches(0.2),
             Inches(0.8), Inches(0.7), size=26)
    add_text(sl, title, bx + Inches(1.1), by + Inches(0.2),
             fw - Inches(1.3), Inches(0.55), size=15, bold=True, color=NAVY)
    add_text(sl, desc, bx + Inches(1.1), by + Inches(0.75),
             fw - Inches(1.3), Inches(1.2), size=12, color=DGRAY)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 6 — 동작 흐름
# ══════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
add_rect(sl, 0, 0, W, H, fill=NAVY)
add_rect(sl, 0, 0, W, Inches(0.08), fill=ACCENT)

add_text(sl, "동작 흐름", Inches(0.55), Inches(0.22), Inches(11), Inches(0.7),
         size=28, bold=True, color=WHITE)
add_text(sl, "SN 입력 한 번으로 전체 파이프라인이 자동 실행됩니다",
         Inches(0.55), Inches(0.85), Inches(11), Inches(0.4),
         size=15, color=LBLUE)
add_rect(sl, Inches(0.4), Inches(1.15), W - Inches(0.8), Pt(1), fill=BLUE)

flow = [
    ("💬", "Knox 챗봇\nSN 입력"),
    ("🖥️", "FA 서비스\n서버 수신"),
    ("🔗", "Superset\n데이터 조회"),
    ("📐", "NETA\n네트워크 분석"),
    ("📄", "PDF\n자동 생성"),
    ("📱", "Knox\n결과 전송"),
]

bw = Inches(1.8)
bh = Inches(3.2)
gap = Inches(0.42)
total = len(flow) * bw + (len(flow) - 1) * gap
sx = (W - total) / 2
by = Inches(1.9)

for i, (icon, label) in enumerate(flow):
    bx = sx + i * (bw + gap)
    c = BLUE if i % 2 == 0 else RGBColor(0x1a, 0x50, 0x8b)
    add_rect(sl, bx, by, bw, bh, fill=c)
    add_rect(sl, bx, by, bw, Inches(0.07), fill=ACCENT)

    add_text(sl, icon, bx, by + Inches(0.25),
             bw, Inches(0.85), size=30, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(sl, label, bx, by + Inches(1.1),
             bw, Inches(1.8), size=13, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER)

    if i < len(flow) - 1:
        add_text(sl, "▶", bx + bw, by + Inches(1.2),
                 gap, Inches(0.6), size=18, color=ACCENT, align=PP_ALIGN.CENTER)

add_text(sl, "⏱ 평균 처리 시간: 약 2~3분  (NETA 데이터 포함)",
         Inches(2), Inches(5.5), Inches(9), Inches(0.6),
         size=15, color=LBLUE, align=PP_ALIGN.CENTER, italic=True)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 7 — 자동화 구성
# ══════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
add_rect(sl, 0, 0, W, H, fill=GRAY)
slide_header(sl, "자동화 구성", "서버만 켜두면 모든 과정이 자동으로 실행됩니다")

schedules = [
    ("🕘", "09:00",   "매일",   "자동 로그인",       "Superset 세션 자동 갱신\n핸드폰 생체인증 승인만 필요"),
    ("🕕", "18:00",   "매일",   "자동 파이프라인",    "FA 미결건 메일 다운로드\n→ 쿼리 → PDF → Knox 전송"),
]

for i, (icon, time, freq, title, desc) in enumerate(schedules):
    by = Inches(1.6) + i * Inches(2.4)
    add_rect(sl, Inches(0.4), by, Inches(12.5), Inches(2.1), fill=WHITE)
    add_rect(sl, Inches(0.4), by, Inches(0.12), Inches(2.1), fill=ACCENT)

    add_text(sl, icon, Inches(0.65), by + Inches(0.4),
             Inches(0.9), Inches(1.0), size=36)
    add_rect(sl, Inches(1.7), by + Inches(0.35),
             Inches(1.8), Inches(1.35), fill=NAVY)
    add_text(sl, time, Inches(1.75), by + Inches(0.38),
             Inches(1.7), Inches(0.7), size=30, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER)
    add_text(sl, freq, Inches(1.75), by + Inches(1.0),
             Inches(1.7), Inches(0.4), size=12, color=LBLUE,
             align=PP_ALIGN.CENTER)
    add_text(sl, title, Inches(3.75), by + Inches(0.25),
             Inches(4), Inches(0.6), size=18, bold=True, color=NAVY)
    add_text(sl, desc, Inches(3.75), by + Inches(0.85),
             Inches(8.5), Inches(1.0), size=13, color=DGRAY)

add_rect(sl, Inches(0.4), Inches(6.15), Inches(12.5), Inches(0.9), fill=NAVY)
add_text(sl, "✅  서버 PC만 켜두면 퇴근 후에도 자동으로 동작합니다",
         Inches(1.0), Inches(6.2), Inches(11), Inches(0.7),
         size=15, bold=True, color=WHITE)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 8 — 기대 효과
# ══════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
add_rect(sl, 0, 0, W, H, fill=GRAY)
slide_header(sl, "기대 효과", "업무 효율이 크게 향상됩니다")

before_after = [
    ("수동 이메일 확인", "자동 다운로드 · 분석"),
    ("SN별 직접 조회",   "챗봇 입력 한 번으로 완료"),
    ("수작업 리포트 작성", "PDF 자동 생성 · 전송"),
    ("결과 전달 지연",     "분석 완료 즉시 Knox 수신"),
]

bw = Inches(5.5)
bh = Inches(1.1)
lx = Inches(0.5)
rx = Inches(7.0)

add_rect(sl, lx, Inches(1.35), bw, Inches(0.45), fill=RGBColor(0xfe, 0xe2, 0xe2))
add_text(sl, "  ❌  개선 전", lx, Inches(1.38), bw, Inches(0.4),
         size=13, bold=True, color=RGBColor(0xdc, 0x26, 0x26))

add_rect(sl, rx, Inches(1.35), bw, Inches(0.45), fill=RGBColor(0xd1, 0xfa, 0xe5))
add_text(sl, "  ✅  개선 후", rx, Inches(1.38), bw, Inches(0.4),
         size=13, bold=True, color=RGBColor(0x05, 0x96, 0x69))

for i, (before, after) in enumerate(before_after):
    by = Inches(1.85) + i * (bh + Inches(0.15))
    add_rect(sl, lx, by, bw, bh, fill=RGBColor(0xff, 0xf1, 0xf1))
    add_text(sl, before, lx + Inches(0.2), by + Inches(0.25),
             bw - Inches(0.4), Inches(0.65), size=15, color=RGBColor(0x7f, 0x1d, 0x1d))

    add_text(sl, "→", Inches(6.15), by + Inches(0.25),
             Inches(0.7), Inches(0.65), size=22, bold=True,
             color=ACCENT, align=PP_ALIGN.CENTER)

    add_rect(sl, rx, by, bw, bh, fill=RGBColor(0xec, 0xfd, 0xf5))
    add_text(sl, after, rx + Inches(0.2), by + Inches(0.25),
             bw - Inches(0.4), Inches(0.65), size=15,
             bold=True, color=RGBColor(0x06, 0x4e, 0x3b))


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 9 — 마무리
# ══════════════════════════════════════════════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
add_rect(sl, 0, 0, W, H, fill=NAVY)
add_rect(sl, 0, 0, Inches(0.25), H, fill=ACCENT)
add_rect(sl, Inches(9), Inches(2), Inches(5), Inches(5), fill=BLUE)
add_rect(sl, Inches(10), Inches(2.8), Inches(4), Inches(4), fill=RGBColor(0x1a, 0x50, 0x8b))

add_text(sl, "🚀", Inches(0.7), Inches(1.5), Inches(2), Inches(1.4), size=60, color=WHITE)
add_text(sl, "Knox SN 입력 한 번,\n통화품질 분석 끝.",
         Inches(0.7), Inches(2.7), Inches(8.5), Inches(2.0),
         size=38, bold=True, color=WHITE)
add_text(sl, "자동화로 더 스마트하게, 더 빠르게.",
         Inches(0.7), Inches(4.7), Inches(8.5), Inches(0.7),
         size=18, color=LBLUE, italic=True)

add_text(sl, "감사합니다",
         Inches(0.7), Inches(6.3), Inches(4), Inches(0.6),
         size=20, bold=True, color=WHITE)

# ── 저장 ──────────────────────────────────────────────────────────────────────
import os as _os
out = _os.path.join(_os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))), "knox_chatbot_소개.pptx")
prs.save(out)
print(f"저장 완료: {out}")
