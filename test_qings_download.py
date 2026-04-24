"""
Qings 엑셀 다운로드 테스트 스크립트

실행:
  python test_qings_download.py

동작:
  1. Qings 사이트 접속 (SSO)
  2. 날짜 / 필터 설정
  3. Apply → 엑셀 다운로드
  4. 저장 경로 출력
"""
import asyncio
import sys
import os

# 프로젝트 루트를 sys.path에 추가
sys.path.insert(0, os.path.dirname(__file__))

# Windows ProactorEventLoop 설정
if sys.platform == "win32":
    asyncio.set_event_loop_policy(asyncio.WindowsProactorEventLoopPolicy())

import logging
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
)

from app.config import settings


async def main():
    print("=" * 60)
    print("Qings 엑셀 다운로드 테스트")
    print("=" * 60)
    print(f"대상 URL  : https://{settings.QINGS_URL}")
    print(f"저장 경로  : {settings.CSV_DOWNLOAD_PATH}")
    print(f"Headless  : {settings.QINGS_HEADLESS}")
    print()

    # QINGS_HEADLESS가 True면 강제로 False로 변경 (SSO 때문)
    if settings.QINGS_HEADLESS:
        print("※ SSO 로그인을 위해 headful 모드로 실행합니다.")
        settings.QINGS_HEADLESS = False

    from app.prefetch.qings_scraper import scrape_qings_excel

    print("스크래핑 시작...\n")
    result = await scrape_qings_excel(save_dir=settings.CSV_DOWNLOAD_PATH)

    print()
    print("=" * 60)
    if result:
        print(f"✅ 다운로드 성공!")
        print(f"   저장 경로: {result}")
        size = os.path.getsize(result)
        print(f"   파일 크기: {size:,} bytes ({size/1024:.1f} KB)")

        # 엑셀 헤더 미리보기
        try:
            import openpyxl
            wb = openpyxl.load_workbook(result, read_only=True, data_only=True)
            ws = wb.active
            headers = [str(c.value).strip() if c.value else "" for c in next(ws.iter_rows(max_row=1))]
            row_count = ws.max_row - 1
            wb.close()
            print(f"   행 수    : {row_count}행")
            print(f"   헤더 목록:")
            for i, h in enumerate(headers):
                print(f"     [{i}] {h}")
        except Exception as e:
            print(f"   (헤더 미리보기 실패: {e})")
    else:
        print("❌ 다운로드 실패 — 로그를 확인하세요.")
    print("=" * 60)


if __name__ == "__main__":
    asyncio.run(main())
